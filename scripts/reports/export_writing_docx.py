"""Convert the frozen writing dossier and appendices into an editable Word report."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
PACK=ROOT/'deliverables/writing_materials_20261024'
OUT=ROOT/'deliverables/Nguyen_lieu_viet_bai_DeFi_20261024.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.85)
sec.header_distance=sec.footer_distance=Inches(.35)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
    s=styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.line_spacing=1.12
styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',22),('Subtitle',11),('Heading 1',15),('Heading 2',13),('Heading 3',11.5)]:
    styles[name].font.size=Pt(size)
    styles[name].paragraph_format.space_before=Pt(13 if name.startswith('Heading') else 0)
    styles[name].paragraph_format.space_after=Pt(7)
    if name.startswith('Heading'):styles[name].font.bold=True
for s in styles:
    if s.type==1:
        # No inherited ornamental title borders or theme colors.
        ppr=s.element.find(qn('w:pPr'))
        if ppr is not None:
            for e in list(ppr):
                if e.tag==qn('w:pBdr'):ppr.remove(e)

def run(p,value,bold=False,italic=False,code=False):
    r=p.add_run(value);r.bold=bold;r.italic=italic
    if code:r.font.name='Consolas';r.font.size=Pt(9)
    return r

def math_run(value):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r

def native_math(p,value):
    fractions={
      'HF_i = sum_j(C_ij × P_j × LT_ij) / risk_debt_i':('HFᵢ = ','Σⱼ Cᵢⱼ × Pⱼ × LTᵢⱼ','risk debtᵢ'),
      'eligible_debt_after_shock / total_debt_after_shock':('','eligible debt after shock','total debt after shock'),
      'p_i=w_i/sum w':('pᵢ = ','wᵢ','Σⱼ wⱼ'),
      '(HHI-1/N)/(1-1/N)':('','HHI − 1/N','1 − 1/N'),
      'E[Y]=q mu/(1-NB(0))':('E[Y] = ','qμ','1 − NB(0)')}
    linear={'HHI=sum p_i²':'HHI = Σᵢ pᵢ²','q_t=logistic(X_t beta)':'qₜ = logistic(Xₜβ)',
      'NB(mu_t,r)':'NB(μₜ, r)','mu_t=softplus(X_t gamma)':'μₜ = softplus(Xₜγ)',
      'P(Y=0)=1-q':'P(Y = 0) = 1 − q','P(Y=k>0)=q NB(k)/(1-NB(0))':'P(Y = k > 0) = q NB(k)/(1 − NB(0))'}
    if value not in fractions and value not in linear:return False
    eq=OxmlElement('m:oMath')
    if value in fractions:
        prefix,num,den=fractions[value]
        if prefix:eq.append(math_run(prefix))
        f=OxmlElement('m:f');n=OxmlElement('m:num');n.append(math_run(num));d=OxmlElement('m:den');d.append(math_run(den));f.append(n);f.append(d);eq.append(f)
    else:eq.append(math_run(linear[value]))
    p._p.append(eq);return True

def rich(p,s):
    pattern=r'(\[([^\]]+)\]\(([^)]+)\)|\*\*(.+?)\*\*|`([^`]+)`|(?<!\*)\*([^*]+)\*(?!\*))'
    last=0
    for m in re.finditer(pattern,s):
        run(p,s[last:m.start()])
        if m.group(2):
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m.group(3),RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'174B72');pr.append(c);r.append(pr)
            t=OxmlElement('w:t');t.text=m.group(2);r.append(t);h.append(r);p._p.append(h)
        elif m.group(4):run(p,re.sub(r'(?<!\w)_([^_]+)_(?!\w)',r'\1',m.group(4)),bold=True)
        elif m.group(5):
            if not native_math(p,m.group(5)):run(p,m.group(5),code=True)
        elif m.group(6):run(p,m.group(6),italic=True)
        last=m.end()
    run(p,s[last:])

def table(lines):
    rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c) for c in r)]
    n=len(rows[0]);t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if n==2:width=[1.5,5.3]
    elif n==3:width=[1.1,2.8,2.9]
    elif n==5:width=[1.8,1.2,1.35,1.2,1.25]
    else:width=[6.8/n]*n
    for col,w in zip(t.columns,width):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate(rows):
        row=t.add_row();trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(width[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if i==0 else ('F6F7F8' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.04
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(value)>45 else WD_ALIGN_PARAGRAPH.CENTER
            rich(p,value)
            for r in p.runs:r.font.size=Pt(10);r.bold=True if i==0 else r.bold
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def heading(s):
    return re.sub(r'[^\w\s]',' ',s,flags=re.UNICODE).strip()

def convert(path,appendix=None):
    lines=path.read_text(encoding='utf8').splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            table(block);continue
        match=re.match(r'^(#{1,3})\s+(.+)',line)
        if match:
            level=len(match[1]);title=match[2]
            if level==1:
                if appendix:doc.add_heading(heading(appendix),level=1)
                else:doc.add_paragraph('Bộ nguyên liệu viết bài về DeFi lending',style='Title')
            else:doc.add_heading(heading(title),level=min(level-1 if not appendix else level,3))
            i+=1;continue
        bullet=re.match(r'^[-*]\s+(.+)',line);number=re.match(r'^\d+\.\s+(.+)',line)
        p=doc.add_paragraph(style='List Bullet' if bullet else 'List Number' if number else 'Normal')
        rich(p,bullet[1] if bullet else number[1] if number else line);i+=1

header=sec.header.paragraphs[0];header.text='NGUYÊN LIỆU NGHIÊN CỨU DEFI LENDING';header.style='Normal'
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Trang ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

convert(PACK/'NGUYEN_LIEU_VIET_BAI.md')
doc.add_page_break();convert(PACK/'BANG_KET_QUA_H1.md','Phụ lục A Kết quả H1 theo 13 quý')
doc.add_page_break();convert(PACK/'TAI_LIEU_THAM_KHAO.md','Phụ lục B Tài liệu tham khảo')
doc.core_properties.title='Bộ nguyên liệu viết bài về DeFi lending'
doc.core_properties.subject='Findings phương pháp dữ liệu kết quả và tài liệu tham khảo H1 H2 H3 và RQ1 đến RQ7'
doc.core_properties.author='Nhóm nghiên cứu DeFi lending'
doc.save(OUT)
print(OUT)
