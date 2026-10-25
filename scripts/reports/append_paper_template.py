"""Append the requested five-part manuscript framework while retaining the existing Word body."""
from pathlib import Path
import re,shutil
from docx import Document

root=Path(__file__).resolve().parents[2]
source=root/'deliverables/Nguyen_lieu_viet_bai_DeFi_20261024.docx'
output=root/'deliverables/Nguyen_lieu_viet_bai_DeFi_20261024_co_template.docx'
backup=root/'output/word_template_20261024/original.docx'
backup.parent.mkdir(parents=True,exist_ok=True)
if not backup.exists():shutil.copy2(source,backup)
doc=Document(source)
doc.add_page_break()
doc.add_heading('Phụ lục C Khung cố định để viết bài',level=1)

def rich(p,line):
    last=0
    for m in re.finditer(r'\*\*(.+?)\*\*',line):
        p.add_run(line[last:m.start()]);p.add_run(m[1]).bold=True;last=m.end()
    p.add_run(line[last:])

lines=(root/'deliverables/TEMPLATE_VIET_BAI_DEFI.md').read_text(encoding='utf8').splitlines()
for line in lines[1:]:
    line=line.strip()
    if not line:continue
    match=re.match(r'^(#{2,3})\s+(.+)',line)
    if match:
        title=re.sub(r'[^\w\s]',' ',match[2],flags=re.UNICODE)
        doc.add_heading(title,level=len(match[1]));continue
    rich(doc.add_paragraph(),line)
doc.save(output)
print(output)
